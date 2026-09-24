#!/usr/bin/env python3
"""Vector Legal — 15 theme variants. Theme engine + content + builder."""
import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

ASSETS = '/tmp/vl_assets'
OUTDIR = os.path.expanduser('~/projects/vector-legal-decks15')
os.makedirs(OUTDIR, exist_ok=True)

SW, SH = Inches(13.333), Inches(7.5)
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

# ---------------------------------------------------------------- themes
def T(name, mode, bg, surface, card, card_line, text, muted, accent, accent2,
      f_title, f_body, f_mono, radius=0.09, art=None, emblem=None,
      kicker_accent=None, rule=True, title_align='center'):
    return dict(name=name, mode=mode, bg=bg, surface=surface, card=card,
        card_line=card_line, text=text, muted=muted, accent=accent, accent2=accent2,
        f_title=f_title, f_body=f_body, f_mono=f_mono, radius=radius, art=art,
        emblem=emblem, kicker_accent=kicker_accent or accent, rule=rule,
        title_align=title_align)

THEMES = [
 T('01-obsidian-neon', 'dark', '0B1220', '111C2E', '141F33', '24344D',
   'F2F7FD', '8FA3BC', '38B2F8', 'B48CF8', 'Inter Display', 'Inter', 'JetBrains Mono',
   art=('full', 'title_hero.png'),
   emblem=('scales_icon_t.png', '38B2F8')),
 T('02-academic-light', 'light', 'F7F8FC', 'FFFFFF', 'FFFFFF', 'DFE4F0',
   '1B2340', '5A6480', '4C6FFF', '8B5CF6', 'Inter Display', 'Inter', 'JetBrains Mono',
   art=('half', 'hero_academic.png')),
 T('03-editorial-cream', 'light', 'FAF6EF', 'FFFFFF', 'FFFDF8', 'E4DCCB',
   '2B241A', '7A6E58', 'A6842E', '6B8F71', 'Playfair Display', 'Inter', 'JetBrains Mono',
   title_align='left'),
 T('04-mono-swiss', 'light', 'FFFFFF', 'FFFFFF', 'FFFFFF', '111111',
   '111111', '555555', 'D92B2B', '111111', 'Inter Display', 'Inter', 'JetBrains Mono',
   radius=0.0, title_align='left'),
 T('05-midnight-terminal', 'dark', '0A0E14', '0E141D', '111826', '1E2A3A',
   'D7E4F0', '6B7E93', '5EEAD4', '4C9EEB', 'Inter Display', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_terminal.png'), emblem=('scales_icon_t.png', '5EEAD4')),
 T('06-paper-minimal', 'light', 'FBFAF7', 'FFFFFF', 'FFFFFF', 'E8E5DE',
   '26241F', '8A857A', 'C4553B', '26241F', 'Noto Serif Display', 'Inter', 'JetBrains Mono',
   title_align='left'),
 T('07-corporate-navy', 'dark', '0F1B2D', '14243B', '162841', '27405F',
   'F5EFE0', 'A9B4C6', 'C9A227', 'E4D5A5', 'Playfair Display', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_navy_gold.png')),
 T('08-vanta-glass', 'dark', '050505', '0C0C0E', '101014', '232330',
   'F4F6F8', '8B929E', '34D399', '9CA8FF', 'Inter Display', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_vanta.png'), emblem=('scales_icon_t.png', '34D399')),
 T('09-light-blueprint', 'light', 'EEF2F7', 'FFFFFF', 'FFFFFF', 'C9D4E4',
   '1C2A3A', '5B6B80', '2F6FED', '0FA3A3', 'Inter Display', 'Inter', 'JetBrains Mono',
   radius=0.04, title_align='left'),
 T('10-violet-grad', 'dark', '12081F', '1A0E2E', '1E1236', '35215A',
   'F5EEFF', 'A793C9', 'E879F9', '8B5CF6', 'Unbounded', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_violet.png'), emblem=('scales_icon_t.png', 'E879F9')),
 T('11-sepia-legal', 'light', 'F4EDE0', 'FBF6EC', 'FBF6EC', 'DECDB2',
   '3A2E1E', '82725A', '7A2E2E', '4A3728', 'Playfair Display', 'Inter', 'JetBrains Mono',
   title_align='left'),
 T('12-arctic-frost', 'light', 'F2F7FA', 'FFFFFF', 'FFFFFF', 'D5E2EA',
   '20303C', '64798A', '3D7EA6', '6BB2C9', 'Inter Display', 'Inter', 'JetBrains Mono'),
 T('13-carbon-orange', 'dark', '161616', '1E1E1E', '202020', '333333',
   'F5F2EE', '9A948C', 'FF6B35', 'FFB38A', 'Unbounded', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_carbon.png'), emblem=('scales_icon_t.png', 'FF6B35'),
   radius=0.03),
 T('14-forest-institute', 'dark', '0E2A1E', '133526', '153A2B', '26523E',
   'F2EBDD', '9DB3A4', 'B8A268', 'D9C896', 'Playfair Display', 'Inter', 'JetBrains Mono',
   art=('full', 'hero_forest.png')),
 T('15-typographic-bw', 'light', 'FFFFFF', 'FFFFFF', 'FFFFFF', '111111',
   '0A0A0A', '666666', '1D4ED8', '0A0A0A', 'Inter Display', 'Inter', 'JetBrains Mono',
   radius=0.0, rule=False, title_align='left'),
]

# ---------------------------------------------------------------- helpers
def _rgb(h): return RGBColor.from_string(h)

def set_alpha(shape, pct):
    """pct = opacity percent of the fill color."""
    s = shape.fill.fore_color._xFill.find('.//' + qn('a:srgbClr'))
    s.append(s.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))}))

def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, radius=None,
             alpha=None, gradient=None, grad_angle=90):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if (radius or 0) > 0 else MSO_SHAPE.RECTANGLE
    sp = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and radius > 0:
        try: sp.adjustments[0] = min(0.5, radius / min(w, h))
        except Exception: pass
    if gradient:
        sp.fill.gradient()
        stops = sp.fill.gradient_stops
        stops[0].color.rgb = _rgb(gradient[0]); stops[0].position = 0.0
        stops[1].color.rgb = _rgb(gradient[1]); stops[1].position = 1.0
        try: sp.fill.gradient_angle = grad_angle
        except Exception: pass
    elif fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = _rgb(fill)
        if alpha is not None: set_alpha(sp, alpha)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = _rgb(line); sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp

def add_text(slide, x, y, w, h, runs, size=12, color='FFFFFF', bold=False,
             font='Inter', align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, line_spacing=None, italic=False):
    """runs: str or list of (text, dict-overrides)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing: p.line_spacing = line_spacing
    if isinstance(runs, str): runs = [(runs, {})]
    for txt, ov in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(ov.get('size', size))
        r.font.bold = ov.get('bold', bold)
        r.font.italic = ov.get('italic', italic)
        r.font.name = ov.get('font', font)
        r.font.color.rgb = _rgb(ov.get('color', color))
        spc = ov.get('spc', spacing)
        if spc: r.font._rPr.set('spc', str(spc))
    return tb

def add_bullets(slide, x, y, w, h, items, size=11.5, color='8FA3BC', font='Inter',
                line_spacing=1.25, gap_after=4, marker='—  ', marker_color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap_after)
        if marker:
            r = p.add_run(); r.text = marker
            r.font.size = Pt(size); r.font.name = font
            r.font.color.rgb = _rgb(marker_color or color)
        r = p.add_run(); r.text = it
        r.font.size = Pt(size); r.font.name = font
        r.font.color.rgb = _rgb(color)
    return tb

def kicker(slide, theme, text, x=0.62, y=0.55):
    add_text(slide, x, y, 6.5, 0.3, text.upper(), size=11, bold=True,
             font=theme['f_mono'], color=theme['kicker_accent'], spacing=200)

def title_block(slide, theme, text, y=0.88, size=37, x=0.62, w=12.1,
                align=None, underline=True):
    al = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER}[align or theme['title_align']]
    if al == PP_ALIGN.CENTER: x = 0.62
    add_text(slide, x, y, w, 0.75, text, size=size, bold=True,
             font=theme['f_title'], color=theme['text'], align=al)
    if underline and theme['rule']:
        uw = 1.15
        ux = x if al == PP_ALIGN.LEFT else 6.67 - uw / 2 + 0.0
        if al == PP_ALIGN.CENTER: ux = (13.333 - uw) / 2
        add_rect(slide, ux, y + 0.72, uw, 0.045, fill=theme['accent'])

# Колонтитул «N / M · бренд». По умолчанию — своя дека Vector Legal; чужая дека
# (deck_builder) обязана выставить своё: иначе в неё уезжают «/ 12» и
# «Vector Legal» (так случилось с vector-prediction: 13 слайдов, «9 / 12»).
FOOTER = {'total': 12, 'brand': 'Vector Legal · Hermes Agent · Osmosy'}

def set_footer(total, brand):
    """Число слайдов и бренд для колонтитулов следующей собираемой деки."""
    FOOTER['total'], FOOTER['brand'] = int(total), str(brand)

def footer(slide, theme, idx, total=None, brand=None):
    c = theme['muted']
    total = total or FOOTER['total']
    add_text(slide, 0.62, 7.14, 1.2, 0.25, f'{idx} / {total}', size=9.5, font=theme['f_mono'], color=c)
    # правый край тот же (12.7"), бокс шире — длинный бренд не переносится
    add_text(slide, 8.2, 7.14, 4.5, 0.25, brand or FOOTER['brand'],
             size=9.5, font=theme['f_mono'], color=c, align=PP_ALIGN.RIGHT)

def bg_fill(slide, theme, decor=True):
    if decor and theme['mode'] == 'dark':
        add_rect(slide, 0, 0, 13.334, 7.5,
                 gradient=(theme['bg'], _shift(theme['bg'], 14)), grad_angle=115)
    else:
        add_rect(slide, 0, 0, 13.334, 7.5, fill=theme['bg'])
    # top hairline accent
    add_rect(slide, 0, 0, 13.334, 0.05, fill=theme['accent'])

def _shift(hexcol, amt):
    r = int(hexcol[0:2], 16) + amt; g = int(hexcol[2:4], 16) + amt; b = int(hexcol[4:6], 16) + amt
    clamp = lambda v: max(0, min(255, v))
    return f'{clamp(r):02X}{clamp(g):02X}{clamp(b):02X}'

def tint_emblem(theme):
    """Recolor transparent emblem to theme accent; returns cached path."""
    src, hexc = theme['emblem']
    out = f'/tmp/vl_assets/emblem_{theme["name"]}.png'
    if os.path.exists(out): return out
    im = Image.open(f'{ASSETS}/{src}').convert('RGBA')
    a = im.getchannel('A')
    solid = Image.new('RGBA', im.size, _rgb(hexc) + (255,))
    solid.putalpha(a)
    solid.save(out)
    return out

def full_art(slide, theme, img, overlay_pct=30):
    pic = slide.shapes.add_picture(f'{ASSETS}/{img}', 0, 0, width=SW, height=SH)
    ov = add_rect(slide, 0, 0, 13.334, 7.5, fill=theme['bg'])
    set_alpha(ov, overlay_pct)
    # raise pic+overlay above bg rect: move both right after bg (positions 2,3,4)
    spTree = slide.shapes._spTree
    for el in [ov._element, pic._element][::-1]:
        spTree.remove(el); spTree.insert(3, el)
    # re-add top accent line above overlay
    add_rect(slide, 0, 0, 13.334, 0.05, fill=theme['accent'])

def half_art(slide, theme, img):
    pic = slide.shapes.add_picture(f'{ASSETS}/{img}', Inches(7.1), 0, width=Inches(6.233), height=SH)
    spTree = slide.shapes._spTree
    spTree.remove(pic._element); spTree.insert(3, pic._element)
    add_rect(slide, 7.1, 0, 0.035, 7.5, fill=theme['accent'])

def card(slide, theme, x, y, w, h, head, lines, head_size=15, body_size=11.5,
         head_color=None, num=None):
    add_rect(slide, x, y, w, h, fill=theme['card'], line=theme['card_line'],
             line_w=1.0, radius=theme['radius'])
    pad = 0.28
    cx = x + pad
    if num is not None:
        d = 0.52
        cxc = x + w / 2 - d / 2
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cxc), Inches(y + 0.26), Inches(d), Inches(d))
        circ.fill.solid(); circ.fill.fore_color.rgb = _rgb(theme['accent'])
        circ.line.fill.background(); circ.shadow.inherit = False
        tf = circ.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(num)
        r.font.size = Pt(16); r.font.bold = True
        r.font.name = theme['f_mono']
        r.font.color.rgb = _rgb('FFFFFF' if theme['mode'] == 'dark' else 'FFFFFF')
        ty = y + 0.92
    else:
        ty = y + 0.24
    add_text(slide, cx, ty, w - 2 * pad, 0.35, head, size=head_size, bold=True,
             font=theme['f_body'], color=head_color or theme['accent'])
    if lines:
        yy = ty + 0.42
        for ln in lines:
            emphasis = ln.startswith('!')
            txt = ln[1:] if emphasis else ln
            add_text(slide, cx, yy, w - 2 * pad, 0.3, txt, size=body_size,
                     font=theme['f_body'],
                     color=theme['text'] if emphasis else theme['muted'])
            yy += 0.285 if body_size <= 11.5 else 0.33

def chip(slide, theme, x, y, w, big, small):
    add_rect(slide, x, y, w, 1.02, fill=theme['card'], line=theme['card_line'],
             line_w=1.0, radius=theme['radius'])
    add_text(slide, x, y + 0.12, w, 0.5, big, size=24, bold=True,
             font=theme['f_title'], color=theme['accent'], align=PP_ALIGN.CENTER)
    add_text(slide, x, y + 0.63, w, 0.3, small, size=10.5, font=theme['f_mono'],
             color=theme['muted'], align=PP_ALIGN.CENTER, spacing=100)

def wide_panel(slide, theme, x, y, w, h, head, lines, head_color=None):
    add_rect(slide, x, y, w, h, fill=theme['surface'], line=theme['card_line'],
             line_w=1.0, radius=theme['radius'])
    add_rect(slide, x, y, 0.05, h, fill=theme['accent'])
    add_text(slide, x + 0.32, y + 0.24, w - 0.7, 0.35, head, size=14.5, bold=True,
             font=theme['f_body'], color=head_color or theme['accent'])
    yy = y + 0.68
    for ln in lines:
        emphasis = ln.startswith('!')
        txt = ln[1:] if emphasis else ln
        add_text(slide, x + 0.32, yy, w - 0.7, 0.3, txt, size=11.5,
                 font=theme['f_body'],
                 color=theme['text'] if emphasis else theme['muted'])
        yy += 0.30

# ---------------------------------------------------------------- slides
def s_title(slide, th):
    full_art(slide, th, th['art'][1], overlay_pct=26) if th['art'] and th['art'][0] == 'full' else None
    if th['art'] and th['art'][0] == 'half':
        half_art(slide, th, th['art'][1])
    ta = th['title_align']
    if ta == 'center' and not (th['art'] and th['art'][0] == 'half'):
        if th['emblem']:
            e = tint_emblem(th)
            slide.shapes.add_picture(e, Inches(6.17), Inches(0.34), Inches(1.0), Inches(1.0))
        pill_w = 3.6
        add_rect(slide, (13.333 - pill_w) / 2, 1.52, pill_w, 0.42,
                 fill=th['surface'], line=th['card_line'], radius=0.21, alpha=70)
        add_text(slide, (13.333 - pill_w) / 2, 1.60, pill_w, 0.3, 'ЮРИДИЧЕСКИЙ AI-ДЕПАРТАМЕНТ',
                 size=10.5, bold=True, font=th['f_mono'], color=th['text'],
                 align=PP_ALIGN.CENTER, spacing=180)
        add_text(slide, 1.67, 2.02, 10.0, 1.2, 'Vector Legal', size=64, bold=True,
                 font=th['f_title'], color=th['accent'], align=PP_ALIGN.CENTER)
        add_text(slide, 3.17, 3.30, 7.0, 0.4, 'Навыки российского права для Hermes Agent',
                 size=17, font=th['f_body'], color=th['text'], align=PP_ALIGN.CENTER)
        for i, (big, small) in enumerate([('12', 'плагинов'), ('167', 'навыков'), ('РФ', 'право')]):
            chip(slide, th, 3.47 + i * 2.25, 4.12, 1.95, big, small)
        add_rect(slide, 4.87, 5.62, 3.6, 0.52, fill=th['accent'], radius=0.26)
        add_text(slide, 4.87, 5.74, 3.6, 0.3, 'github.com/Osmosy/vector-legal', size=12.5,
                 bold=True, font=th['f_mono'],
                 color='FFFFFF' if th['mode'] == 'dark' else 'FFFFFF', align=PP_ALIGN.CENTER)
        add_text(slide, 3.67, 6.55, 6.0, 0.28, 'Адаптация из anthropics/claude-for-legal · Apache-2.0',
                 size=10.5, font=th['f_body'], color=th['muted'], align=PP_ALIGN.CENTER)
        add_text(slide, 4.87, 6.90, 3.6, 0.26, 'Osmosy · Hermes Agent · 2026', size=9.5,
                 font=th['f_mono'], color=th['muted'], align=PP_ALIGN.CENTER)
    else:
        # left-aligned title (light editorial / half-art themes)
        add_rect(slide, 0.62, 1.30, 0.05, 2.2, fill=th['accent'])
        add_text(slide, 0.92, 1.30, 6.5, 0.3, 'ЮРИДИЧЕСКИЙ AI-ДЕПАРТАМЕНТ', size=11,
                 bold=True, font=th['f_mono'], color=th['kicker_accent'], spacing=200)
        add_text(slide, 0.88, 1.66, 6.2, 1.15, 'Vector Legal', size=58, bold=True,
                 font=th['f_title'], color=th['accent'])
        add_text(slide, 0.92, 2.86, 5.9, 0.7, 'Навыки российского права для Hermes Agent',
                 size=16.5, font=th['f_body'], color=th['text'])
        for i, (big, small) in enumerate([('12', 'плагинов'), ('167', 'навыков'), ('РФ', 'право')]):
            x = 0.92 + i * 1.95
            add_text(slide, x, 3.85, 1.8, 0.5, big, size=26, bold=True,
                     font=th['f_title'], color=th['accent'])
            add_text(slide, x, 4.36, 1.8, 0.3, small, size=10, font=th['f_mono'],
                     color=th['muted'], spacing=100)
        add_text(slide, 0.92, 5.15, 6.0, 0.3, 'github.com/Osmosy/vector-legal', size=13,
                 bold=True, font=th['f_mono'], color=th['accent'])
        add_text(slide, 0.92, 5.55, 6.0, 0.28, 'Адаптация из anthropics/claude-for-legal · Apache-2.0',
                 size=10.5, font=th['f_body'], color=th['muted'])
        add_text(slide, 0.92, 5.90, 6.0, 0.26, 'Osmosy · Hermes Agent · 2026', size=9.5,
                 font=th['f_mono'], color=th['muted'])

def s_intro(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '01 · Введение')
    title_block(slide, th, 'Что такое Vector Legal')
    add_text(slide, 0.62, 1.86, 11.0, 0.35, 'Набор юридических навыков для Hermes Agent',
             size=15.5, bold=True, font=th['f_body'], color=th['accent'])
    add_text(slide, 0.62, 2.30, 12.0, 0.6,
             'Аналогия: как нанять ассистента, который знает внутренние правила вашей команды,\nпороги эскалаций и стиль писем — по вашим правилам, не generic-шаблонам.',
             size=12, font=th['f_body'], color=th['muted'], line_spacing=1.15)
    cards = [
        ('Профиль практики', ['Агент работает по внутренним правилам', 'вашей команды, извлечённым из', 'реальных подписанных договоров']),
        ('12 направлений', ['Коммерция · ПДн · корпоратив ·', 'трудовое · суды · ИС · ИИ-гавернанс', '!Все — право Российской Федерации']),
        ('Валидация человеком', ['Отправку наружу, подписания делает', 'человек. Агент только готовит драфт', 'и помогает с ревизией']),
    ]
    for i, (h, l) in enumerate(cards):
        card(slide, th, 0.62 + i * 4.11, 3.16, 3.75, 1.55, h, l)
    wide_panel(slide, th, 0.62, 5.02, 12.1, 1.85, 'Из чего выросло', [
        'Основа: anthropics/claude-for-legal (Anthropic, Apache-2.0) — 151 навык в 12 плагинах',
        'Адаптация (Osmosy, Apache-2.0) — русское право + Hermes Agent + расширения:',
        '!протоколы разногласий вместо redline · kad.arbitr/pravo.gov.ru вместо CourtListener · ЕГРЮЛ/ФИПС вместо SEC/USPTO · 4-ступенчатый контракт для cron-агентов',
    ])
    footer(slide, th, 2)

def s_workflow(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '02 · Как это работает')
    title_block(slide, th, '4 шага от установки до продакшена')
    steps = [
        ('Установка', ['git clone + подключение', 'в Hermes Agent']),
        ('Cold-start интервью', ['2–15 минут · 10–92 вопроса', 'Извлекает playbook команды']),
        ('Работа навыков', ['Ревизия / триаж / драфт', 'по внутренним правилам']),
        ('Мониторинг', ['5 cron-агентов', 'НПА · продления · споры']),
    ]
    xs = [0.55, 3.80, 7.05, 10.30]
    for i, (h, l) in enumerate(steps):
        card(slide, th, xs[i], 2.10, 2.92, 2.05, h, l, num=i + 1)
        if i < 3:
            add_rect(slide, xs[i] + 2.97, 3.05, 0.28, 0.035, fill=th['accent'])
    wide_panel(slide, th, 0.83, 4.85, 11.67, 1.95, 'Режим «Без настройки»', [
        'До подключения practice profile агент работает с дефолтами РФ (средний риск-аппетит, типовые позиции по законам).',
        'Каждый выход помечается меткой — «требует настройки профиля».',
        'Для калибровки — cold-start interview. Юрист отвечает 2–15 минут.',
    ])
    footer(slide, th, 3)

DOMAINS1 = [
    ('Коммерческие договоры', ['Ревизия входящих договоров против playbook, NDA-триаж', 'ГК 15/330/401/425/452 · ч. 5 ст. 4 АПК']),
    ('Персональные данные', ['152-ФЗ: оператор/поручитель, запросы субъектов, утечки', 'ст. 6/9/12/14/18.1/21 — ТИПЗ ФСТЭК — КоАП 13.11']),
    ('Корпоративное право', ['Одобрения, протоколы 181.2 ГК, DD dataroom, M&A с эскроу', 'ФЗ-14 ст. 38/45/46 · ФЗ-208 ст. 79/84 · ФЗ-135 ФАС']),
    ('Трудовое право', ['Найм → увольнение → ЛНА. Ат-вилл не работает (закрытый перечень)', 'ТК 15/57/70/77-83/180/193/373 — Пленум ВС №15']),
    ('Судебная работа', ['Досудебный порядок ч. 5 ст. 4 АПК, kad.arbitr, claim-chart', 'АПК 4/66/72/88/125/131/259/276 — Пленумы ВС РФ']),
    ('ИИ-гавернанс', ['Реестр ИИ-систем, AIA, вендорские ИИ-оговорки, shadow-AI', 'ЭПР ФЗ-258 [verify] — ГОСТ Р 59276 — 152-ФЗ — 187-ФЗ КИИ']),
]
def s_domains1(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '03 · Домены · 1/2')
    title_block(slide, th, '6 основных направлений')
    for i, (h, l) in enumerate(DOMAINS1):
        x = 0.62 if i % 2 == 0 else 6.88
        y = 1.88 + (i // 2) * 1.66
        card(slide, th, x, y, 5.83, 1.46, h, l)
    footer(slide, th, 4)

def s_domains2(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '03 · Домены · 2/2')
    title_block(slide, th, '6 дополнительных направлений')
    cards = [
        ('Регуляторный мониторинг', ['pravo.gov.ru, гильотина ПП №1128, 44-ФЗ/223-ФЗ, ФАС', 'Отраслевые: Минсельхоз / Россельхознадзор / Минпромторг']),
        ('Интеллектуальная собственность', ['Clearance по ФИПС, takedown ст. 1253.1, OSS-совместимость', 'ГК ч. 4: 1225-1254 · 1354-1375 · СИПН · патентный поверенный']),
        ('Продукты и лицензии', ['Launch-review, реклама ФЗ-38 + ЕРИР-маркировка, ЗоЗПП', 'Оферта 437/428 ГК + политика 152-ФЗ + ФЗ-161/115/54']),
        ('Обучение студентов', ['LEARNING MODE — Socratic, IRAC-анализ, case brief по ВС РФ', 'Экзамен на статус адвоката ФЗ-63 гл. 2-3']),
        ('Юридические клиники', ['ФЗ-324, супервизор проверяет каждое письмо клиенту', 'Согласие 152-ФЗ + intake + семестровая отчётность']),
        ('Каталог навыков', ['Установка, QA 740 строк (13 параметров), SHA-pinning', 'Injection-scan при установке каждого чужого навыка']),
    ]
    for i, (h, l) in enumerate(cards):
        x = 0.62 if i % 2 == 0 else 6.88
        y = 1.88 + (i // 2) * 1.66
        card(slide, th, x, y, 5.83, 1.46, h, l)
    add_rect(slide, 0.62, 6.82, 12.1, 0.035, fill=th['accent'])
    add_text(slide, 0.62, 6.94, 12.0, 0.25,
             'Расширения: vector-check (DD контрагента) · privacy-policy-ru · terms-of-service-ru · legal-research-ru · patent-claim-chart',
             size=11, font=th['f_body'], color=th['muted'])
    footer(slide, th, 5)

def s_architecture(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '04 · Архитектура')
    title_block(slide, th, 'Профиль практики — ядро всего')
    add_text(slide, 0.62, 1.84, 11.5, 0.3,
             'Каждый навык работает от practice profile. Без него — метка «требует настройки».',
             size=12.5, font=th['f_body'], color=th['text'])
    wide_panel(slide, th, 0.83, 2.35, 11.67, 2.05, 'Файл практики — ~/.hermes/legal/<домен>/CLAUDE.md', [
        'Внутренние правила: лимит ответственности, indemnity, 152-ФЗ, сроки, суд',
        'Матрица эскалаций с денежными порогами',
        'Стиль работы, реальные договоры как основа, главная позиция',
        'Плейсхолдеры до cold-start — все навыки работают в режиме «требует настройки»',
    ])
    card(slide, th, 0.83, 4.72, 5.62, 2.15, 'Общие правила — в каждом навыке', [
        'Замечание для ревизора — один блок над документом',
        'Дальше — варианты развития',
        'Проверка на конфиденциальность содержимого',
        'Указание источника цитаты, объём, актуальность нормы',
    ], body_size=11)
    card(slide, th, 6.88, 4.72, 5.62, 2.15, 'Валидация человеком', [
        'Удаление контента / претензия / подписание — только драфт',
        'Отправка — человек (никогда агент)',
        'Патентные поверенные — обязательны (ФЗ-316)',
        'Заявки Роспатент, споры СИПН — intake, дальше специалист',
    ], body_size=11)
    footer(slide, th, 6)

def s_agents(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '05 · Агенты мониторинга')
    title_block(slide, th, '5 автономных агентов по расписанию')
    cards = [
        ('Монитор НПА', ['Понедельник 09:00 — pravo.gov.ru, regulation.gov.ru', 'Проверка материальности, разрывы, еженедельный digest', 'Направление: регуляторный мониторинг']),
        ('Монитор продлений', ['Понедельник 09:00 — договоры, дедлайн cancel-by', 'Дата отправки с учётом бизнес-дней РФ', 'Направление: коммерческие договоры']),
        ('Монитор судебных дел', ['Ежедневно 08:00 — kad.arbitr.ru, sudrf.ru, ФИПС', 'Новые подачи, дедлайны, постановления', 'Если дедлайн меньше 14 дней — уведомление в Telegram']),
        ('Монитор запусков продуктов', ['Понедельник 10:00 — продуктовый трекер', '9-категорийный фреймворк — риск-классификация', 'Направление: продуктовый юрист']),
        ('Таблица DD dataroom', ['Под конкретную сделку — пакетная обработка VDR', 'corporate-legal']),
        ('Модель безопасности', ['Только чтение до leaf · пишущий лист один — для результата', 'No external send без подтверждения · lead, не вывод']),
    ]
    for i, (h, l) in enumerate(cards):
        x = 0.83 if i % 2 == 0 else 6.88
        y = 1.98 + (i // 2) * 1.72
        hh = 1.5 if i < 4 else 1.25
        card(slide, th, x, y, 5.62, hh, h, l, body_size=11)
    footer(slide, th, 7)

def s_mcp(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '06 · Подключение правовых баз')
    title_block(slide, th, 'MCP-коннекторы к РФ-базам')
    add_text(slide, 0.62, 1.86, 12.0, 0.3,
             'По умолчанию навыки работают через web_search с указанием источника.',
             size=12, font=th['f_body'], color=th['muted'])
    add_text(slide, 0.62, 2.20, 12.0, 0.3,
             'MCP-серверы снимают лишний флаг — данные сразу приходят с provenance автоматически.',
             size=12, font=th['f_body'], color=th['muted'])
    add_text(slide, 0.62, 2.72, 4.0, 0.3, 'ДОСТУПНЫ СЕГОДНЯ', size=11, bold=True,
             font=th['f_mono'], color=th['kicker_accent'], spacing=180)
    cards = [
        ('atomno-mcp', ['13 серверов, Python, MIT', 'ЕГРЮЛ / ЕГРИП — через ФНС открытые данные', 'Контрагент одним вызовом — ЕФРСБ, Картотека, ФССП', 'sudact — судебная практика РФ'], 0.83, 3.15, 5.96, 1.95),
        ('Russian-Law-MCP', ['npm @ansvar/russian-law-mcp · Apache-2.0', '12 369 федеральных законов · 77 647 положений', '«КонсультантПлюс, но для AI»', 'Базовая правовая база для всех 12 доменов'], 6.88, 3.15, 5.62, 1.95),
        ('ГАРАНТ MCP (официальный API)', ['Подключение через токен Гарант-Коннект', 'Полные редакции · судебная практика (Сутяжник)', 'Мониторинг изменений — currency trigger по-русски'], 0.83, 5.32, 5.96, 1.55),
        ('Резервный режим: web_search + [verify]', ['Без MCP все цитаты — [модель знания — проверь]', 'или прямая проверка против pravo.gov.ru / ФИПС / ЕГРЮЛ'], 6.88, 5.32, 5.62, 1.55),
    ]
    for h, l, x, y, w, hh in cards:
        card(slide, th, x, y, w, hh, h, l, body_size=11)
    footer(slide, th, 8)

def s_connectors(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '07 · Распределение коннекторов')
    title_block(slide, th, 'Какая база для какого домена')
    rows = [
        ('Коммерческие договоры', 'Russian-Law · ЕГРЮЛ · + Контрагент 1-выз. · закупки · ЕРИР'),
        ('Персональные данные', 'Russian-Law (152-ФЗ) · + Контрагент-1'),
        ('Корпоративные сделки', 'ЕГРЮЛ · Russian-Law · + ЦБ-ставки · Росреестр'),
        ('Трудовое право', 'Russian-Law (ТК РФ) · + ФНС-калькулятор (взносы, НДФЛ)'),
        ('Судебная работа', 'sudact · Контрагент-1 · + ЦБ-ставки (ст. 395 проценты)'),
        ('ИИ-гавернанс', 'Russian-Law (152-ФЗ) · + реестр ЭПР ИИ — в работе'),
        ('Регуляторный мониторинг', 'Russian-Law · ЦБ-ставки · pravo.gov.ru (в roadmap MCP)'),
        ('Интеллектуальная собственность', 'sudact · Роспатент/ФИПС товарные знаки · Росреестр'),
    ]
    for i, (h, l) in enumerate(rows):
        x = 0.62 if i % 2 == 0 else 6.88
        y = 1.95 + (i // 2) * 1.12
        add_rect(slide, x, y, 5.83, 0.94, fill=th['card'], line=th['card_line'],
                 line_w=1.0, radius=th['radius'])
        add_text(slide, x + 0.26, y + 0.14, 5.3, 0.3, h, size=13, bold=True,
                 font=th['f_body'], color=th['accent'])
        add_text(slide, x + 0.26, y + 0.50, 5.4, 0.3, l, size=10.5, font=th['f_body'],
                 color=th['muted'])
    add_text(slide, 0.62, 6.55, 12.0, 0.28,
             'Продукты: ЕРИР-реклама · Честный ЗНАК (roadmap) · клиники/каталог: ЕГРЮЛ + sudact',
             size=11, font=th['f_body'], color=th['muted'])
    footer(slide, th, 9)

def s_guide(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '08 · Руководство пользователя')
    title_block(slide, th, 'Что делать и как')
    steps = [
        ('Выбрать домен', ['12 плагинов — не ставь все сразу,', 'поставь 1-2 по своей роли', '→ Установка за 5 минут']),
        ('Пройти интервью', ['Cold-start: 2 мин (quick)', 'или 15 мин (full)', '→ Профиль заполняется']),
        ('Использовать', ['Все навыки домена работают', 'по правилам твоей команды', '→ Выход — draft для проверки юристом']),
    ]
    for i, (h, l) in enumerate(steps):
        card(slide, th, 0.83 + i * 3.96, 1.95, 3.75, 2.2, h, l, num=i + 1, body_size=11)
    wide_panel(slide, th, 0.83, 4.55, 11.67, 2.25, 'Типовые запросы к агенту', [
        '«Проверь этот договор поставки на риски» → Коммерческие договоры',
        '«Нужна ли оценка воздействия для этой фичи?» → Персональные данные',
        '«Кто-то использует наш логотип без разрешения» → ИС',
        '«Сделай хронологию по делу» → Судебная работа',
    ])
    footer(slide, th, 10)

def s_roadmap(slide, th):
    bg_fill(slide, th)
    kicker(slide, th, '09 · Развитие')
    title_block(slide, th, 'Roadmap')
    items = [
        ('Один MCP-сервер vector-legal', ['Единый сервер поверх существующих коннекторов: atomno-mcp + Russian-Law-MCP', 'Инструменты «проверь контрагента», «найди судебную практику», «загрузи редакцию»', 'Снимает требование [verify] — provenance приходит автоматически'], 1.90, 1.62),
        ('Отраслевые надстройки', ['Мясопереработка · сельское хозяйство · ФЗ-193 о сельскохозяйственной кооперации', 'Гос-контракты 44-ФЗ/223-ФЗ · Антимонопольный домен ФЗ-135', 'Расширяет commercial-legal отраслевыми правилами'], 3.72, 1.62),
        ('Связки с экосистемой Vector', ['AgentHub · Work / Meat / Marketing — общий practice profile между доменами Vector'], 5.54, 1.10),
    ]
    for h, l, y, hh in items:
        wide_panel(slide, th, 0.83, y, 11.67, hh, h, l)
    footer(slide, th, 11)

def s_final(slide, th):
    if th['art'] and th['art'][0] == 'full':
        full_art(slide, th, th['art'][1], overlay_pct=42)
    else:
        bg_fill(slide, th)
    add_text(slide, 2.67, 1.35, 8.0, 1.0, 'Vector Legal', size=52, bold=True,
             font=th['f_title'], color=th['accent'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 2.50, 8.0, 0.35,
             '12 плагинов · 167 навыков · адаптация claude-for-legal под право РФ',
             size=14, font=th['f_body'], color=th['text'], align=PP_ALIGN.CENTER)
    add_rect(slide, 4.17, 3.05, 5.0, 0.035, fill=th['accent'])
    add_text(slide, 4.17, 3.55, 5.0, 0.32, 'github.com/Osmosy/vector-legal', size=14.5,
             bold=True, font=th['f_mono'], color=th['text'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 4.25, 8.0, 0.28,
             'Apache-2.0 · структурная база: anthropics/claude-for-legal (Anthropic)',
             size=11.5, font=th['f_body'], color=th['muted'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 4.58, 8.0, 0.28, 'адаптация под закон России: Osmosy',
             size=11.5, font=th['f_body'], color=th['muted'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 5.40, 8.0, 0.3,
             'Запустить: git clone https://github.com/Osmosy/vector-legal.git',
             size=12, font=th['f_body'], color=th['text'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 5.95, 8.0, 0.3, '«Пройти cold-start интервью commercial-legal»',
             size=13, bold=True, font=th['f_body'], color=th['accent'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 6.30, 8.0, 0.26, '— этого достаточно, чтобы начать работу.',
             size=10.5, font=th['f_body'], color=th['muted'], align=PP_ALIGN.CENTER)
    add_text(slide, 2.67, 6.95, 8.0, 0.24,
             'Все выходные данные — черновики для проверки юристом. Не юридическая консультация.',
             size=9.5, font=th['f_body'], color=th['muted'], align=PP_ALIGN.CENTER)

BUILDERS = [s_title, s_intro, s_workflow, s_domains1, s_domains2, s_architecture,
            s_agents, s_mcp, s_connectors, s_guide, s_roadmap, s_final]

def build_theme(th):
    set_footer(len(BUILDERS), 'Vector Legal · Hermes Agent · Osmosy')
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    blank = prs.slide_layouts[6]
    for builder in BUILDERS:
        slide = prs.slides.add_slide(blank)
        builder(slide, th)
    out = f'{OUTDIR}/vector-legal-{th["name"]}.pptx'
    prs.save(out)
    return out

if __name__ == '__main__':
    import sys
    only = sys.argv[1:] if len(sys.argv) > 1 else None
    for th in THEMES:
        if only and th['name'] not in only: continue
        p = build_theme(th)
        print('built', p)
