#!/usr/bin/env python3
"""Дымовая сборка: все темы Vector Legal + референс-дека deck-music.py.

Hero-арты живут вне репо, поэтому каждый PNG, на который ссылаются engine.py и
deck-music.py, подменяется заглушкой во временном VECTOR_DECK_ASSETS. Проверяется
не дизайн, а контракт движка:

- каждая тема собирается без исключений;
- колонтитул «N / M»: M — реальное число слайдов деки;
- бренд колонтитула — свой у каждой деки (дефект, из-за которого тест появился:
  footer() писал «/ 12 · Vector Legal» во все деки, в т.ч. в 13-слайдовую
  vector-prediction).

Запуск:  python3 tests/smoke_build.py   (нужны python-pptx, Pillow)
"""
from __future__ import annotations

import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def placeholder_assets(target: Path) -> None:
    from PIL import Image

    names = set()
    for src in ("engine.py", "deck-music.py"):
        names |= set(re.findall(r"['\"]([A-Za-z0-9_\-]+\.png)['\"]", (ROOT / src).read_text(encoding="utf-8")))
    # Имена из f-строк вида f'{ASSETS}/vector_ray_t.png' — отдельный проход.
    for src in ("engine.py", "deck-music.py"):
        names |= set(re.findall(r"\}/([A-Za-z0-9_\-]+\.png)", (ROOT / src).read_text(encoding="utf-8")))
    for name in names:
        # Логотип Vector Ray — в реальном размере 1300×920: тест `check_logo`
        # проверяет его наличие по габаритам, заглушка 640×360 его не заменит.
        size = (1300, 920) if name == "vector_ray_t.png" else (640, 360)
        Image.new("RGBA", size, (40, 60, 90, 255)).save(target / name)


def slide_texts(pptx: Path) -> list[list[str]]:
    with zipfile.ZipFile(pptx) as z:
        slides = sorted(
            (int(m.group(1)), n) for n in z.namelist()
            if (m := re.match(r"ppt/slides/slide(\d+)\.xml$", n))
        )
        return [re.findall(r"<a:t>([^<]*)</a:t>", z.read(n).decode("utf-8")) for _, n in slides]


def check_footers(pptx: Path, brand: str) -> list[str]:
    errors = []
    slides = slide_texts(pptx)
    total = len(slides)
    for i, texts in enumerate(slides, 1):
        for t in texts:
            m = re.fullmatch(r"(\d+) / (\d+)", t)
            if m and (int(m.group(1)) != i or int(m.group(2)) != total):
                errors.append(f"{pptx.name} слайд {i}: колонтитул «{t}», ждали «{i} / {total}»")
            if "Hermes Agent · Osmosy" in t and t != brand:
                errors.append(f"{pptx.name} слайд {i}: бренд «{t}», ждали «{brand}»")
    return errors


def slide_service_colors(pptx: Path) -> list[tuple[int, str, str]]:
    """Цвета служебных строк в собранном PPTX: (слайд, текст, цвет).

    Берём цвет из XML, а не из кода движка: проверка должна падать, если
    `footer()`/`signature_line()` начнут писать muted, даже когда сам
    `service_color()` остался правильным (мутация «вернуть muted» ловилась
    только чтением собранного файла).
    """
    out: list[tuple[int, str, str]] = []
    with zipfile.ZipFile(pptx) as z:
        slides = sorted(
            (int(m.group(1)), n) for n in z.namelist()
            if (m := re.match(r"ppt/slides/slide(\d+)\.xml$", n))
        )
        for num, name in slides:
            xml = z.read(name).decode("utf-8")
            for para in re.findall(r"<a:p>.*?</a:p>", xml, re.S):
                texts = re.findall(r"<a:t>([^<]*)</a:t>", para)
                if not texts:
                    continue
                text = "".join(texts)
                if not (re.fullmatch(r"\d+ / \d+", text) or "Osmosy · Hermes Agent" in text):
                    continue
                m = re.search(r'<a:solidFill><a:srgbClr val="([0-9A-Fa-f]{6})"', para)
                if m:
                    out.append((num, text, m.group(1).upper()))
    return out


def check_credit_contrast(eng, pptx: Path | None = None, theme: dict | None = None) -> list[str]:
    """Служебные строки читаются: контраст >= 4.5 (WCAG AA) на фоне темы.

    Дефект (27.09.2026): на пяти тёмных темах с full-артом строка лежала на
    арте цветом muted — контраст 1.14–1.42, почти не читалась. Цвет выбирает
    движок (eng.credit_color), подложка возвращает фон темы, поэтому контраст
    считается от bg — это нижняя граница того, что видит читатель.

    Задача C1: тем же цветом muted рисовались колонтитул (footer) и подпись
    «Osmosy · Hermes Agent · 2026» — на светлых темах 06/11/12 контраст 3.52/
    4.00/4.19 при норме 4.5. Если переданы `pptx` и `theme`, цвета служебных
    строк берутся из собранного слайда (тогда мутация «вернуть muted в
    footer()» падает), иначе — по коду движка.
    """
    def hx(s):
        return tuple(int(s[i:i+2], 16) for i in (0, 2, 4))

    errors = []
    for th in eng.THEMES:
        c = eng._contrast(hx(eng.credit_color(th)), hx(th['bg']))
        if c < 4.5:
            errors.append(f"тема {th['name']}: строка атрибуции — контраст {c:.2f} < 4.5")
        sc = eng._contrast(hx(eng.service_color(th)), hx(th['bg']))
        if sc < 4.5:
            errors.append(f"тема {th['name']}: служебные строки — контраст {sc:.2f} < 4.5")

    if pptx is not None and theme is not None:
        for num, text, color in slide_service_colors(pptx):
            c = eng._contrast(hx(color), hx(theme['bg']))
            if c < 4.5:
                errors.append(
                    f"{pptx.name} слайд {num}: служебная строка «{text}» цвет {color} "
                    f"— контраст {c:.2f} < 4.5 (тема {theme['name']})"
                )
    return errors


def slide_logo_pictures(pptx: Path, min_side: int = 900) -> list[tuple[int, int]]:
    """Крупные картинки на титуле — для проверки логотипа Vector Ray.

    Дефект (27.09.2026): при пересборке музыкальных дек из `deck-music.py`
    логотип Vector Ray (1300×920) пропал с титула — он был только в старых
    pptx, а исходник его не рисовал. Файлы упали с 1.35 МБ до 0.14 МБ.
    Порог 900: у логотипа меньшая сторона 920, у hero-арта 1152 — оба видны,
    а мелкие элементы (эмблема 1″ ≈ 300 px) не попадают.
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    out: list[tuple[int, int]] = []
    prs = Presentation(str(pptx))
    slide = prs.slides[0]
    for sh in slide.shapes:
        if sh.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        w, h = sh.image.size
        if min(w, h) >= min_side:
            out.append((w, h))
    return out


def check_logo(pptx: Path) -> list[str]:
    """На титуле деки есть логотип Vector Ray (1300×920)."""
    pics = slide_logo_pictures(pptx)
    if not any((w, h) == (1300, 920) for w, h in pics):
        return [f"{pptx.name}: на титуле нет логотипа Vector Ray (1300×920); "
                f"картинки: {pics or 'нет'}"]
    return []


def main() -> int:
    with tempfile.TemporaryDirectory() as tmp:
        assets, out = Path(tmp, "assets"), Path(tmp, "out")
        assets.mkdir()
        placeholder_assets(assets)
        os.environ["VECTOR_DECK_ASSETS"] = str(assets)
        os.environ["VECTOR_DECK_OUTDIR"] = str(out)
        sys.path.insert(0, str(ROOT))

        import deck_builder  # грузит engine.py рядом с собой, с путями из окружения
        eng = deck_builder.eng

        errors: list[str] = []
        errors += check_credit_contrast(eng)
        for th in eng.THEMES:
            try:
                p = Path(eng.build_theme(th))
            except Exception as e:  # noqa: BLE001 — нужна вся картина, а не первое падение
                errors.append(f"тема {th['name']}: {type(e).__name__}: {e}")
                continue
            if len(slide_texts(p)) != len(eng.BUILDERS):
                errors.append(f"{p.name}: слайдов {len(slide_texts(p))}, ждали {len(eng.BUILDERS)}")
            errors += check_footers(p, "Vector Legal · Hermes Agent · Osmosy")
            errors += check_credit_contrast(eng, p, th)

        for p in map(Path, deck_builder.build(str(ROOT / "deck-music.py"))):
            if Path(p).parent != out:
                errors.append(f"{p}: VECTOR_DECK_OUTDIR не учтён")
            errors += check_footers(p, "Vector Music · Hermes Agent · Osmosy")
            errors += check_logo(p)
            if any("Vector Legal" in t for s in slide_texts(p) for t in s):
                errors.append(f"{p.name}: в чужой деке текст «Vector Legal»")

    for e in errors:
        print(f"ERROR {e}")
    print(f"\nИтог: тем {len(eng.THEMES)}, ошибок {len(errors)}")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
